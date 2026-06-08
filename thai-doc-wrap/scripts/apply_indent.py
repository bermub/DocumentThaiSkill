#!/usr/bin/env python3
"""
apply_indent.py — ตั้งย่อหน้าบรรทัดแรก (first-line indent) และระยะแท็บ
ให้เป็นมาตรฐานสากลตามรูปแบบงานวิจัย

มาตรฐานสากล (APA 7th / MLA 9th / Chicago / Turabian):
  - ย่อหน้าบรรทัดแรกของทุกย่อหน้าเนื้อความ = 0.5 นิ้ว = 1.27 ซม. = 720 DXA
  - ระยะแท็บเริ่มต้น (กด Tab 1 ครั้ง) = 0.5 นิ้ว เท่ากับย่อหน้าพอดี
  - (มาตรฐานวิทยานิพนธ์ มจพ. ใช้ 1.0 ซม. — เลือกด้วย --cm 1.0)

ตั้งย่อหน้าให้เฉพาะ "ย่อหน้าเนื้อความ" เท่านั้น โดยข้าม:
  หัวข้อ (Heading/Title/TOC/Caption) ย่อหน้าที่จัดกึ่งกลาง/ชิดขวา และเซลล์ตาราง
"""
import argparse
import sys

DEFAULT_CM = 1.27  # 0.5 นิ้ว — มาตรฐานสากล
SKIP_STYLE_PREFIX = ("Heading", "Title", "TOC", "Caption", "Subtitle")


def _twips(cm: float) -> int:
    return round(cm / 2.54 * 1440)


def _set_default_tab(doc, twips: int):
    from docx.oxml.ns import qn
    settings = doc.settings.element
    el = settings.find(qn("w:defaultTabStop"))
    if el is None:
        el = settings.makeelement(qn("w:defaultTabStop"), {})
        settings.append(el)
    el.set(qn("w:val"), str(twips))


def _is_body(p) -> bool:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    style = (p.style.name or "") if p.style else ""
    if style.startswith(SKIP_STYLE_PREFIX):
        return False
    if p.alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
        return False
    if not p.text.strip():
        return False
    return True


def apply_indent(in_path: str, out_path: str, cm: float = DEFAULT_CM,
                 include_tables: bool = False) -> int:
    from docx import Document
    from docx.shared import Cm

    doc = Document(in_path)
    tw = _twips(cm)
    _set_default_tab(doc, tw)

    n = 0
    paras = list(doc.paragraphs)
    if include_tables:
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    paras.extend(cell.paragraphs)

    for p in paras:
        if _is_body(p):
            p.paragraph_format.first_line_indent = Cm(cm)
            n += 1

    doc.save(out_path)
    return n


def main():
    ap = argparse.ArgumentParser(
        description="ตั้งย่อหน้าบรรทัดแรกมาตรฐานสากลตามงานวิจัย (.docx)")
    ap.add_argument("input", help="ไฟล์ .docx เข้า")
    ap.add_argument("-o", "--output", help="ไฟล์ออก (ค่าเริ่มต้น *.indent.docx)")
    ap.add_argument("--cm", type=float, default=DEFAULT_CM,
                    help="ระยะย่อหน้า (ซม.) ค่าเริ่มต้น 1.27 (สากล) | มจพ.=1.0")
    ap.add_argument("--tables", action="store_true",
                    help="ย่อหน้าในเซลล์ตารางด้วย (ปกติงานวิจัยไม่ย่อในตาราง)")
    args = ap.parse_args()

    out = args.output or args.input.rsplit(".", 1)[0] + ".indent.docx"
    n = apply_indent(args.input, out, cm=args.cm, include_tables=args.tables)
    print(f"✓ ตั้งย่อหน้า {args.cm} ซม. ({_twips(args.cm)} DXA) ให้ {n} ย่อหน้า → {out}",
          file=sys.stderr)


if __name__ == "__main__":
    main()
